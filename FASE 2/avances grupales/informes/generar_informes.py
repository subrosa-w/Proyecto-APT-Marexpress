# -*- coding: utf-8 -*-
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path(__file__).resolve().parent
DATOS = json.loads((BASE.parent / "backend" / "scripts" / "resultado-admin-operador.json").read_text(encoding="utf-8"))

AZUL = RGBColor(0x0B, 0x3A, 0x6A)


def sombra(celda, color):
    tc = celda._tePr if hasattr(celda, "_tePr") else celda._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def configurar(doc, titulo):
    sec = doc.sections[0]
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.header.paragraphs[0].text = "Marexpress  ·  " + titulo
    sec.footer.paragraphs[0].text = "Informe del 11 de octubre de 2026  ·  uso interno del proyecto"
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    estilo.paragraph_format.space_after = Pt(8)
    for nombre, tam in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        doc.styles[nombre].font.color.rgb = AZUL
        doc.styles[nombre].font.name = "Calibri"
        doc.styles[nombre].font.size = Pt(tam)


def p(doc, texto, negrita=False):
    par = doc.add_paragraph()
    run = par.add_run(texto)
    run.bold = negrita
    return par


def codigo(doc, texto):
    par = doc.add_paragraph()
    run = par.add_run(texto)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    par.paragraph_format.left_indent = Cm(0.4)
    return par


def tabla(doc, headers, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        celda = t.rows[0].cells[i]
        celda.text = ""
        run = celda.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        sombra(celda, "0B3A6A")
    for fila in filas:
        celdas = t.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = str(valor)
            for par in celdas[i].paragraphs:
                for run in par.runs:
                    run.font.size = Pt(9)
    if anchos:
        for fila in t.rows:
            for i, ancho in enumerate(anchos):
                fila.cells[i].width = Cm(ancho)
    doc.add_paragraph()
    return t


def portada(doc, titulo, subtitulo):
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("MAREXPRESS")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = AZUL
    t = doc.add_heading(titulo, 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s = doc.add_paragraph(subtitulo)
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d = doc.add_paragraph("11 de octubre de 2026")
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def informe_pruebas():
    doc = Document()
    configurar(doc, "Informe de pruebas")
    portada(
        doc,
        "Informe de pruebas del sistema",
        "Administrador, operador y peoneta. Resultados reales contra la API y la base MySQL.",
    )
    doc.add_heading("1. Cómo se probó", 1)
    p(doc, "Las pruebas no fueron una revisión solo de pantalla. Se levantó la API de Marexpress en un puerto de prueba, se firmaron sesiones reales de administrador, operador y peoneta con el mismo mecanismo JWT del sistema, y se llamó a las rutas que usa cada vista. La base fue la MySQL local marexpress_db. Cuando una prueba creaba un dato, se anulaba o se borraba al terminar, salvo el manifiesto que ya había salido a ruta: ese registro se conserva porque la salida es un hecho operativo.")
    p(doc, "Hubo dos tandas el mismo día. La primera, de 43 comprobaciones, cubrió salud del servidor, seguridad, peoneta, entrega y rechazo. La segunda, de 65 comprobaciones, recorrió la vista del administrador y después la del operador. En las dos tandas el resultado fue correcto: 108 de 108.")

    doc.add_heading("2. Resumen", 1)
    tabla(doc, ["Tanda", "Enfoque", "Correctas", "Fallos"], [
        ["Primera", "Transversal, seguridad y peoneta", "43", "0 después de corregir la entrega"],
        ["Segunda", "Administrador y operador", "65", "0"],
        ["Total", "Las tres vistas actuales", "108", "0 abiertos"],
    ])
    p(doc, "Durante la primera tanda la entrega del peoneta sí falló. La base exige abrir la entrega en EN_PROCESO y cerrarla solo cuando cada bulto tiene resultado. Ese flujo se corrigió y se volvió a probar: una orden quedó entregada con pago Transbank y otra quedó rechazada con observación. Esas órdenes de prueba se anularon después.")

    doc.add_heading("3. Pruebas del administrador", 1)
    p(doc, "La vista del administrador abre en /administrador. Desde ahí se probaron los datos que alimentan inicio, órdenes, clientes, tarifas, flota, usuarios, manifiestos, bodega, seguimiento, alertas y reportes. Cuenta corriente y facturación siguen siendo pantallas de aviso: no tienen todavía una operación propia, y por eso no se les inventó un POST.")
    filas_admin = [
        [r["metodo"], r["ruta"], str(r["status"] or "—"), "Correcta" if r["ok"] else "Fallo", r["detalle"]]
        for r in DATOS["resultados"] if r["vista"] == "Administrador"
    ]
    tabla(doc, ["Método", "Ruta", "Código", "Resultado", "Qué respondió"], filas_admin)

    doc.add_heading("Qué se comprobó en cada pantalla", 2)
    p(doc, "Inicio. Pide el informe del mes. Respondió 200 tanto el periodo mensual como el diario del 11 de octubre de 2026. Ese JSON es el que dibujan los gráficos de vueltas, clientes y entregas.")
    p(doc, "Órdenes y seguimiento. El administrador lista órdenes y estados. También anuló la orden de prueba OT-000065. Anular no borra la fila: cambia el estado y conserva el historial, que es lo que la base permite.")
    p(doc, "Clientes, comunas y destinatarios. Las tres lecturas respondieron 200. El operador también puede verlas, porque las usa al crear una orden. El peoneta recibió 403 en clientes.")
    p(doc, "Tarifas. Un tarifario sin nombre fue rechazado con 400. Uno llamado PRUEBA INFORME 2026 se creó (201), se editó el valor por kilo de 100 a 120 (200), se desactivó (200) y su historial guardó 4 cambios. Después se eliminó de la base para no dejar un tarifario falso. El operador intentó crear tarifario y tarifa de trayecto: ambos devolvieron 403.")
    p(doc, "Conductores y vehículos. Un RUT 1-9 y una patente 123 fueron rechazados con el mensaje de formato. No se dejó flota falsa en esta tanda.")
    p(doc, "Usuarios y roles. El listado de usuarios y el catálogo de roles respondieron 200. El catálogo llega envuelto como { roles: [...] }, que es justo lo que lee la pantalla. En una prueba anterior del mismo día, cambiar el rol de un operador a su mismo rol respondió 200. El código impide quitar el rol al último administrador activo.")
    p(doc, "Manifiestos, bodega y alertas. El administrador puede leerlos. La campana de alertas también es suya y del operador.")

    doc.add_heading("4. Pruebas del operador", 1)
    p(doc, "La vista del operador abre en /operador. Se recorrió recepción, nueva orden, listado, clientes, bodega, manifiestos, pronto despacho, seguimiento, pagos y alertas.")
    filas_op = [
        [r["metodo"], r["ruta"], str(r["status"] or "—"), "Correcta" if r["ok"] else "Fallo", r["detalle"]]
        for r in DATOS["resultados"] if r["vista"] == "Operador"
    ]
    tabla(doc, ["Método", "Ruta", "Código", "Resultado", "Qué respondió"], filas_op)

    doc.add_heading("Recorrido real de una orden", 2)
    p(doc, "1. Calcular una orden vacía devolvió 400: falta la dirección de origen. Esa es la misma validación que ve el formulario de Nueva OT.")
    p(doc, "2. Calcular una orden completa, con cliente, comunas, zona urbana, un bulto y pago por pagar, devolvió 200.")
    p(doc, "3. Crear esa orden devolvió 201 y el número OT-000065.")
    p(doc, "4. La recepción buscó ese número con registrar=1 y devolvió 200. Ese parámetro anota en el seguimiento que el documento fue escaneado.")
    p(doc, "5. El pago de la OT se consultó en 200. El estado se cambió a PROGRAMADA y después se devolvió al estado original. La edición de la referencia respondió 200.")
    p(doc, "6. Se creó el manifiesto MAN-000015, se asignó al peoneta de prueba, se escaneó el bulto OT-000065-001 y el manifiesto salió a ruta. Las tres operaciones respondieron bien.")
    p(doc, "7. El administrador anuló OT-000065. El manifiesto MAN-000015 queda como evidencia de una salida a ruta de prueba.")

    doc.add_heading("5. Aislamiento entre vistas", 1)
    filas_ais = [
        [r["nombre"], r["metodo"], str(r["status"] or "—"), r["detalle"]]
        for r in DATOS["resultados"] if r["vista"] in ("Aislamiento", "Seguridad")
    ]
    tabla(doc, ["Prueba", "Método", "Código", "Mensaje"], filas_ais)
    p(doc, "El operador no abre reportes, usuarios ni roles, y no crea tarifarios. El peoneta no abre reportes, usuarios, el listado general de órdenes ni clientes. Un texto con comilla y OR 1=1 en la dirección de una orden fue bloqueado con 400 y quedó registrado como intento de intrusión, con la IP 127.0.0.1 y la ruta POST /api/ordenes-transporte/calcular.")

    doc.add_heading("6. Tanda anterior: peoneta y base del sistema", 1)
    p(doc, "Esa tanda se ejecutó antes, el mismo día, y cerró en 43 correctas. Se resume aquí para que este informe junte todo lo probado.")
    tabla(doc, ["Prueba", "Resultado"], [
        ["GET /api/health y la raíz de la API", "200"],
        ["Ruta inexistente", "404 e incluye el método y la ruta"],
        ["Listado de órdenes sin token", "401"],
        ["Login con OR 1=1", "400 y alerta de intrusión"],
        ["15 lecturas: órdenes, clientes, bodega, flota, alertas, reportes y manifiestos del peoneta", "200"],
        ["Peoneta pide todas las órdenes o los reportes", "403"],
        ["Cliente pide usuarios", "403"],
        ["Alta de conductor y de camión, luego limpieza", "201"],
        ["Cambio de rol, cambio de estado, edición de orden y de manifiesto", "200"],
        ["Entrega sin observación de rechazo", "400, pide el comentario"],
        ["Entrega con receptor y Transbank", "200"],
        ["Segunda entrega de la misma orden", "400, ya tiene resultado"],
        ["Entrega de una OT que no es suya", "404"],
        ["Rechazo con observación", "200"],
        ["DELETE de la orden de prueba", "200, queda anulada"],
        ["Marcar una alerta como leída y devolverla a no leída", "200"],
    ])
    p(doc, "El fallo que apareció y se corrigió antes de dar por buena esa tanda fue este: al guardar la entrega directo como ENTREGADA, MySQL respondía que una entrega nueva debe comenzar en EN_PROCESO. Al marcar un rechazo con cantidad cero, otro control de la base rechazaba entrega_carga porque la cantidad entregada tiene que ser mayor que cero. El servicio ahora abre la entrega, marca los bultos y recién entonces cierra el resultado. En un rechazo no crea fila de cantidad.")

    doc.add_heading("7. Lo que estas pruebas no tapan", 1)
    p(doc, "Cuenta corriente y facturación, dentro del menú del administrador, muestran una página de aviso. No hay todavía un alta ni un listado propio detrás de esos dos títulos.")
    p(doc, "El portal del cliente autentica y no deja ver las otras vistas, pero sus pantallas de órdenes, seguimiento, documentos, pagos y evidencias siguen en aviso. La tabla cliente_usuario está pensada para filtrar las órdenes de esa empresa y todavía no se usa en esa vista.")
    p(doc, "El ingreso real del navegador pide, después de la contraseña, el código que llega al correo. Estas pruebas firmaron el token en el servidor para poder recorrer las rutas sin depender del correo. No se desactivó esa verificación.")
    p(doc, "Para usar en el navegador el cierre de entrega del peoneta hay que reiniciar el proceso de la API si ese proceso se levantó antes de la corrección.")
    p(doc, "Quedaron datos de prueba identificables: órdenes anuladas desde OT-000059 hasta OT-000065, y el manifiesto MAN-000015 en salida a ruta. En una tanda anterior también quedaron el conductor de RUT 22.222.222-2 y el camión ZX99, porque ya estaban ligados a un manifiesto y la base no deja borrarlos. El conductor 33.333.333-3 y el camión ZX98 se crearon y se eliminaron.")
    p(doc, "Hay una copia antigua de páginas en frontend/src/pages. La aplicación que se abre usa frontend/src/sistema. Las pruebas corresponden a esa carpeta.")

    doc.add_heading("8. Conclusión", 1)
    p(doc, "Administrador, operador y peoneta responden en sus rutas, rechazan lo que no les corresponde y validan los datos mal formados antes de guardarlos. El recorrido completo del operador, desde el cálculo hasta la salida a ruta y la anulación, funcionó en esta ejecución. No quedó un error abierto de POST, PUT, PATCH o DELETE en las operaciones probadas.")
    destino = BASE / "Informe-pruebas-Marexpress.docx"
    doc.save(destino)
    return destino


def informe_codigos():
    doc = Document()
    configurar(doc, "Informe de código")
    portada(
        doc,
        "Informe de código",
        "Qué hace cada parte del sistema y cómo se conectan las tres vistas.",
    )

    doc.add_heading("1. De la pantalla a la base", 1)
    p(doc, "El navegador carga React. La aplicación activa está en frontend/src/sistema. App.jsx decide qué pantalla mostrar según la dirección. Si la dirección empieza por /administrador, /operador o /peoneta, primero pasa por ProtectedRoute: sin sesión vuelve a /acceso, y con un rol distinto lo manda a la casa de su rol.")
    p(doc, "Cada pantalla pide datos con frontend/src/sistema/api.js. Esa función agrega el token Bearer, pide al servidor sin usar caché y, si la respuesta trae ruta o código de error, los agrega al mensaje que ve la persona.")
    p(doc, "El servidor es Express. backend/src/app.js recibe el JSON, pasa el filtro de inyección y registra las rutas de backend/src/routes.js. Cada módulo tiene rutas, un controlador y un servicio. El servicio habla con MySQL mediante Prisma. El cliente de Prisma se genera en backend/generated/prisma y se conecta con el adaptador de MariaDB.")
    codigo(doc, "Pantalla React  →  api.js  →  Express  →  autorizarRoles  →  servicio  →  Prisma  →  MySQL")

    doc.add_heading("2. La sesión", 1)
    p(doc, "El login no entrega el token en el primer paso. login.service.js comprueba usuario y contraseña con bcrypt. Si coinciden, verificacionAcceso.service.js crea un desafío y correoOtp.service.js envía un código al correo. Solo cuando ese código es válido, token.service.js firma un JWT de 8 horas. El token lleva el id del usuario, el rol, el emisor marexpress-backend y la audiencia marexpress-web. La clave de firma sale de JWT_SECRET y no se imprime.")
    p(doc, "auth.middleware.js no se fía del rol escrito dentro del token. Después de verificar la firma, vuelve a leer el usuario en la base. Si está desactivado, responde 401. Si el administrador le cambió el rol hace un minuto, el permiso nuevo vale aunque el token siga vigente. El objeto que queda en la petición es { idUsuario, idRol, rol }.")
    p(doc, "Eso importa en el peoneta: el id con el que se buscan los manifiestos sale de ese objeto, nunca del cuerpo que envía el navegador. Una persona no puede pedir los manifiestos de otra cambiando un número en la petición.")

    doc.add_heading("3. Quién puede entrar a cada ruta", 1)
    p(doc, "autorizarRoles se ejecuta después de autenticar. Recibe la lista de roles permitidos. Si el rol de la base no está en esa lista, responde 403 con el texto «No tienes permisos para realizar esta acción». Hay una equivalencia: cuando una ruta dice CLIENTE, también acepta CLIENTE NORMAL y CLIENTE CUENTA CORRIENTE. No hay una equivalencia parecida entre administrador, operador y peoneta. Cada uno entra solo donde su ruta lo nombra.")
    p(doc, "En el navegador hay una segunda puerta. ProtectedRoute compara el rol con la lista de la ruta. RoleLayout tiene el mapa rolesDelPrefijo: admin solo ADMINISTRADOR, operador solo OPERADOR, peoneta solo PEONETA. Si alguien escribe a mano /administrador estando como operador, el layout calcula rolAjeno y, después de ejecutar todos los efectos, navega a /operador. El cálculo va después de los efectos para no romper el orden de los hooks de React.")
    codigo(doc, "rolesDelPrefijo.admin = [\"ADMINISTRADOR\"]\nrolesDelPrefijo.operador = [\"OPERADOR\"]\nrolesDelPrefijo.peoneta = [\"PEONETA\"]")

    doc.add_heading("4. Validación de inyección", 1)
    p(doc, "bloquearInyeccion.middleware.js recorre el cuerpo, la consulta y los parámetros hasta seis niveles de profundidad. Si un texto coincide con un patrón de sentencia SQL, no sigue hacia el controlador. Los patrones cubren UNION SELECT, DROP TABLE, INSERT INTO, DELETE FROM, UPDATE ... SET, information_schema, SLEEP o BENCHMARK, xp_, una comilla seguida de OR, OR 1=1, un punto y coma seguido de un verbo SQL, y comentarios /* */.")
    p(doc, "Antes de responder, llama a registrarAlertaSeguridad. Guarda la IP de la conexión y, si viene distinta, la cabecera X-Forwarded-For. Guarda la ruta y el nombre del campo. Si el campo se llama password, clave, token, codigo o authorization, en la alerta solo dice «dato sensible», para no copiar la contraseña. La respuesta al navegador es 400 y no explica cuál patrón coincidió.")
    p(doc, "Este filtro no reemplaza las consultas parametrizadas. Prisma envía los valores aparte del SQL. El filtro sirve para frenar el intento, anotarlo y avisar. En la prueba, la dirección «calle' OR 1=1 --» nunca llegó a calcular una tarifa.")

    doc.add_heading("5. Cuando algo se cae", 1)
    p(doc, "Si ninguna ruta coincide, app.js responde 404 con el método y la dirección, por ejemplo GET /api/ruta-inexistente. Si un controlador lanza un error no previsto, el manejador final escribe en la consola del servidor la pila y un código ERR- más la hora. Al navegador le manda un mensaje general, la ruta y ese código. No manda la pila. En React, FronteraError envuelve todas las rutas y, si una pantalla se rompe al dibujarse, muestra la dirección y el mensaje del error.")
    p(doc, "Los errores previstos no usan ese camino. Una orden incompleta, un RUT mal escrito o un rol sin permiso responden 400 o 403 con una frase concreta. En la entrega del peoneta, un rechazo de la base se traduce al mensaje del trigger, para que la persona vea la regla y no un texto de Prisma.")

    doc.add_heading("6. La orden de transporte", 1)
    p(doc, "NuevaOrdenPage junta cliente, origen, destino, documento, servicio, cargas y pago. Antes de guardar puede llamar a POST /api/ordenes-transporte/calcular. validarNuevaOrdenTransporte.middleware.js exige dirección de origen y destino, cliente, comunas, zona URBANA o LEJANA, al menos una carga con cantidad y tipo de bulto, y un tipo de documento de la lista FACTURA, BOLETA, GUIA, OT_MANUAL o SIN_DOCUMENTO. Si hay documento, exige número. Si es SIN_DOCUMENTO, no deja poner número.")
    p(doc, "El servicio calcularValoresOrden busca la tarifa del cliente y la tarifa de la zona. El precio base de la carga es opcional: si el formulario lo deja vacío, valorBase se trata como 0 y no se copia solo desde el tarifario, salvo que la petición mande valorBaseAplicado. El total sale de base, kilos, metros cúbicos y zona.")
    p(doc, "Al crear, el servicio abre una transacción: orden, detalle de carga, bultos y pago. El pago exige quién paga, REMITENTE o DESTINATARIO, y el tipo CONTADO, POR_PAGAR o CUENTA_CORRIENTE. En contado el método tiene que ser EFECTIVO, TRANSFERENCIA o TARJETA, y el estado queda PAGADO. Crédito y débito se guardan como TARJETA. Por pagar deja el método vacío y el estado PENDIENTE. Cuenta corriente exige una cuenta. Después, si todo salió bien, se intenta crear la alerta OT_CREADA. Si esa alerta falla, la orden no se deshace.")
    p(doc, "Cambiar el estado es otro camino: PATCH /api/ordenes-transporte/:id/estado. cambiarEstadoOrdenTransporte busca el idEstado en la tabla estado_ot. Si la orden ya está en ese estado, responde 400. Si no, actualiza la orden y escribe historial y seguimiento con registrarEventoOt. La pantalla del operador ya no adivina la bandeja por el manifiesto: bandejaSeguimiento mira primero el nombre guardado en estado_ot.")
    p(doc, "Editar con PUT exige de nuevo origen, destino y comunas, porque el servicio reescribe esos campos. Por eso la pantalla manda el formulario completo, no un solo campo. Anular con DELETE no ejecuta un DELETE de SQL sobre la orden: anularOrdenTransporte cambia el estado a anulada y deja el registro.")

    doc.add_heading("7. Manifiesto, bodega y pago", 1)
    p(doc, "crearManifiesto exige camión, conductor, ciudad, fecha y al menos un peoneta. Crea el manifiesto en PREPARACION, le pone el número MAN- más el id, y llama a asignarPeonetas, que escribe manifiesto_usuario. Escanear un código busca primero el bulto y, si no existe, la OT. Solo carga bultos que estén EN_BODEGA. Sacar a ruta exige camión, conductor, peoneta y al menos un bulto. Si existe el estado EN_TRANSITO, las órdenes del manifiesto pasan a ese estado.")
    p(doc, "La bodega lista bultos por sucursal y estado. Recibir un bulto y reingresar una OT rechazada son POST separados. Pronto despacho lista lo que vence dentro de los días pedidos; la pantalla usa 3.")
    p(doc, "Pagos pendientes y el ambiente de Webpay están abiertos a administrador y operador. Iniciar Webpay crea la transacción con el SDK de Transbank. El retorno de Webpay es público, porque lo llama el proveedor, y queda antes del middleware de autenticación.")

    doc.add_heading("8. La entrega del peoneta, el tramo más exigente", 1)
    p(doc, "obtenerManifiestosPeoneta busca manifiesto_usuario del id del token. De ahí baja al manifiesto, a sus órdenes, a los bultos y al pago más reciente. La pantalla ManifiestosPeonetaPage usa la misma respuesta en tres direcciones: Mi ruta muestra todo, Entregas filtra lo pendiente e Historial filtra lo que ya tiene resultado.")
    p(doc, "registrarEntregaPeoneta solo acepta ENTREGADA o RECHAZADA. Entregada exige nombre de quien recibe y un método entre EFECTIVO, TRANSFERENCIA, TRANSBANK y CHEQUE. Rechazada exige una observación de al menos 3 caracteres. Busca la orden solo si el manifiesto tiene a ese usuario. Si ya hay un resultado final, no deja repetirlo.")
    p(doc, "La base tiene triggers que el código tiene que respetar en este orden, dentro de una transacción:")
    codigo(doc, "1. INSERT entrega_ot con resultadoEntrega = EN_PROCESO\n2. Si es ENTREGADA, INSERT entrega_carga con cantidadEntregada > 0\n3. INSERT entrega_bulto: ENTREGADO o RECHAZADO\n4. UPDATE entrega_ot al resultado final\n5. UPDATE del pago a PAGADO si estaba pendiente y no es cuenta corriente\n6. registrarEventoOt en historial y seguimiento")
    p(doc, "El insert directo como ENTREGADA lo rechaza el trigger trg_entrega_ot_insert. El update a ENTREGADA lo rechaza si la cantidad de entrega_carga no coincide con los bultos marcados ENTREGADO, o si falta algún bulto del manifiesto. Un rechazo no puede tener bultos ENTREGADO, y chk_entrega_carga_cantidad impide guardar cantidad 0, así que en un rechazo no se crea entrega_carga. chk_entrega_motivo obliga a guardar el motivo cuando el resultado es RECHAZADA. Después del update, otro trigger de la base también mueve el estado de la orden a ENTREGADA o RECHAZADA, si ese nombre existe en estado_ot. El servicio crea ese estado si todavía no está.")

    doc.add_heading("9. Alertas, reportes y formatos", 1)
    p(doc, "La tabla alerta guarda tipo, título, mensaje, fecha, si fue leída, y opcionalmente la orden y el usuario. Al crear una OT el tipo es OT_CREADA. Al listar alertas, sincronizarCargaAtrasada compara la fecha de Chile con fechaEntregaEstimada, se salta anuladas, canceladas y entregadas, y deja una alerta no leída por orden atrasada. Si la orden ya no está atrasada, marca la alerta como leída. El tipo INTENTO_INTRUSION sale del filtro de seguridad. AlertasCampana.jsx está en el encabezado del administrador y del operador. Se refresca al cambiar de página y cuando la pestaña vuelve a estar visible. Después de crear una orden, NuevaOrdenPage dispara el evento marexpress-alertas para que la campana se actualice sin recargar.")
    p(doc, "obtenerInforme arma el día o el mes pedido. Si el periodo no es «dia», usa el mes. La fecha, si no tiene forma AAAA-MM-DD, es el día de Chile. Junta órdenes creadas en el rango, manifiestos que salieron y entregas registradas. Separa las órdenes anuladas o canceladas del total activo. ReportesPage exporta CSV con punto y coma y BOM, para Excel, y un PDF horizontal con jsPDF. El inicio del administrador usa el mismo informe para los gráficos de GraficoBarras.")
    p(doc, "formatos.js concentra el formato de los campos: RUT, teléfono, correo, patente, dinero, kilos y centímetros. Las pantallas de login, registro, cliente, destinatario, nueva orden, edición, recepción, manifiestos y tarifas reutilizan esos objetos en el input. validarRut recorre el cuerpo de derecha a izquierda con factores del 2 al 7 y compara el dígito. La patente de un camión nuevo tiene que cumplir dos a cuatro letras y dos números.")

    doc.add_heading("10. Qué archivo corresponde a cada vista", 1)
    tabla(doc, ["Vista", "Pantalla", "Archivo", "Servicio que llama"], [
        ["Administrador", "Inicio y gráficos", "AdminInicioPage.jsx", "GET /api/reportes"],
        ["Administrador", "Reportes y exportación", "ReportesPage.jsx", "GET /api/reportes"],
        ["Administrador", "Usuarios", "UsuariosPage.jsx", "GET usuarios y roles, PATCH rol"],
        ["Administrador", "Flota", "MaestrosPage.jsx", "GET y POST conductores y vehículos"],
        ["Administrador", "Tarifas", "TarifasPage.jsx", "Tarifario, historial, estado y trayecto"],
        ["Administrador", "Órdenes y seguimiento", "OrdenesListadoPage y SeguimientoPage", "Listar, editar, anular, cambiar estado"],
        ["Operador", "Inicio", "OperadorInicioPage.jsx", "Órdenes y bodega de la sucursal"],
        ["Operador", "Recepción", "RecepcionPage.jsx", "OT por número y reingreso"],
        ["Operador", "Nueva OT", "NuevaOrdenPage.jsx", "Calcular y crear"],
        ["Operador", "Manifiestos", "ManifiestosPage.jsx", "Crear, escanear, quitar bulto, salir"],
        ["Operador", "Bodega y despacho", "BodegaPage y ProntoDespachoPage", "Inventario, recibir, pronto despacho"],
        ["Operador", "Pagos", "PagosOtPage.jsx", "Pendientes, consulta y Webpay"],
        ["Peoneta", "Ruta, entregas e historial", "ManifiestosPeonetaPage.jsx", "GET manifiestos y POST entregas"],
    ])
    p(doc, "Las tres vistas comparten RoleLayout, Message, los estilos de sistema.css y, donde corresponde, las tablas mx-table y las tarjetas admin-card o peoneta-card. El peoneta no tiene una hoja de estilos aparte: usa las mismas clases para que la entrega se vea como el resto del sistema.")

    doc.add_heading("11. Qué conviene recordar al leer el código", 1)
    p(doc, "El rol verdadero está en la base, no en el token. El peoneta solo ve manifiestos ligados a su id. La entrega no es un insert único: la base obliga un protocolo. Anular no es borrar. El filtro de SQL avisa y bloquea, y Prisma sigue siendo quien evita que un valor se mezcle con la consulta. Si una pantalla del administrador dice cuenta corriente o facturación, hoy es un aviso, no un módulo terminado.")
    destino = BASE / "Informe-codigos-Marexpress.docx"
    doc.save(destino)
    return destino


def informe_roles():
    doc = Document()
    configurar(doc, "Informe de roles")
    portada(
        doc,
        "Qué hace cada rol",
        "Permisos de administrador, operador y peoneta en pantalla y en la API.",
    )

    doc.add_heading("1. Cómo se decide el permiso", 1)
    p(doc, "El rol vive en la tabla rol y se asigna al usuario. Al entrar, la persona llega a una sola casa: el administrador a /administrador, el operador a /operador y el peoneta a /peoneta. Si escribe la dirección de otra casa, ProtectedRoute y RoleLayout lo devuelven a la suya. En el servidor, cada ruta repite la regla con autorizarRoles. Aunque alguien arme la petición a mano, sin el rol correcto recibe 403.")
    p(doc, "El cliente existe como rol, en las variantes CLIENTE, CLIENTE NORMAL y CLIENTE CUENTA CORRIENTE, pero su portal todavía no muestra sus órdenes. Este informe se concentra en las tres vistas que ya operan: administrador, operador y peoneta.")

    doc.add_heading("2. Administrador", 1)
    p(doc, "El administrador mira la operación completa y además configura lo que los demás usan: personas, roles, tarifas, conductores y camiones. También puede hacer el trabajo operativo, porque varias rutas de órdenes, bodega y manifiestos lo incluyen junto al operador. No reparte: si entra a la ruta del peoneta, el servidor se lo permite por supervisión, pero solo verá manifiestos si él mismo está asignado. En la práctica su casa es el panel administrativo.")
    tabla(doc, ["Pantalla", "Dirección", "Qué puede hacer"], [
        ["Inicio", "/administrador", "Ver gráficos del mes: órdenes, vueltas y entregas"],
        ["Órdenes", "/administrador/ordenes", "Listar, editar, cambiar estado y anular"],
        ["Clientes", "/administrador/clientes", "Ver clientes y comunas, y dar de alta"],
        ["Tarifas", "/administrador/tarifas", "Crear, editar, activar y ver el historial"],
        ["Flota", "/administrador/maestros", "Añadir conductores y camiones"],
        ["Usuarios", "/administrador/usuarios", "Ver cuentas y cambiar el rol"],
        ["Manifiestos", "/administrador/manifiestos", "Ver, crear, cargar y sacar a ruta"],
        ["Bodega", "/administrador/bodega", "Ver inventario y recibir bultos"],
        ["Seguimiento", "/administrador/seguimiento", "Ver el estado y actualizarlo"],
        ["Reportes", "/administrador/reportes", "Informe diario o mensual, Excel y PDF"],
        ["Cuenta corriente", "/administrador/cuentas", "Pantalla de aviso. Aún sin módulo"],
        ["Facturación", "/administrador/facturacion", "Pantalla de aviso. Aún sin módulo"],
    ])
    p(doc, "En la API, solo el administrador puede: listar usuarios, listar roles, cambiar un rol, crear un tarifario, editarlo, activarlo o desactivarlo, ver su historial, crear una tarifa de trayecto, crear conductor y crear vehículo. También es el único que pide GET /api/reportes. No puede dejar la empresa sin ningún administrador activo: cambiarRolUsuario cuenta los administradores y rechaza el cambio si quedaría cero.")
    p(doc, "La campana de alertas y el aviso de actualización también aparecen en su encabezado.")

    doc.add_heading("3. Operador", 1)
    p(doc, "El operador mueve la carga del día. Crea la orden, la recibe, la deja en bodega, arma el manifiesto, lo saca a ruta, sigue el estado y cobra. No configura la empresa: no asigna roles, no crea tarifarios y no exporta el informe de rentabilidad.")
    tabla(doc, ["Pantalla", "Dirección", "Qué puede hacer"], [
        ["Inicio", "/operador", "Resumen de órdenes y bultos de su bodega"],
        ["Recepción", "/operador/recepcion", "Escanear una OT o un bulto y registrar el ingreso"],
        ["Nueva OT", "/operador/ordenes", "Calcular y crear la orden, con alerta incluida"],
        ["Órdenes", "/operador/ordenes-listado", "Buscar, editar, cambiar estado y anular"],
        ["Clientes", "/operador/clientes", "Consultar y crear clientes y destinatarios"],
        ["Bodega", "/operador/bodega", "Inventario, recepción y reingreso de rechazo"],
        ["Manifiestos", "/operador/manifiestos", "Crear, escanear bultos, quitar y salir a ruta"],
        ["Pronto despacho", "/operador/despacho", "Ver lo que vence dentro de 3 días"],
        ["Seguimiento", "/operador/seguimiento", "Ver la bandeja y actualizar el estado"],
        ["Pago de OT", "/operador/pagos", "Ver pendientes, consultar un pago e iniciar Webpay"],
    ])
    p(doc, "Puede leer conductores, camiones, peonetas y tarifas, porque los necesita para armar un manifiesto o calcular una orden. No puede crearlos ni modificar el tarifario: esas pruebas devolvieron 403. No puede abrir /api/reportes, /api/maestros/usuarios ni /api/maestros/roles. Sí ve la campana de alertas, incluida la de carga atrasada y la de intento de intrusión.")

    doc.add_heading("4. Peoneta", 1)
    p(doc, "El peoneta reparte lo que ya le asignaron. No crea órdenes, no arma manifiestos y no ve la flota, los clientes ni los reportes. Su mundo es el manifiesto en el que aparece su usuario.")
    tabla(doc, ["Pantalla", "Dirección", "Qué puede hacer"], [
        ["Mi ruta", "/peoneta", "Ver sus manifiestos, la carga y el destino"],
        ["Entregas", "/peoneta/entregas", "Cerrar lo pendiente: entregada o rechazada"],
        ["Historial", "/peoneta/historial", "Ver entregadas y rechazadas, con comentario y pago"],
        ["Mi perfil", "/peoneta/perfil", "Ver los datos de su cuenta"],
    ])
    p(doc, "GET /api/peoneta/manifiestos devuelve solo las asignaciones de manifiesto_usuario del token. POST /api/peoneta/entregas exige el número de OT, el resultado, y según el caso el receptor, la observación y el método de pago: efectivo, transferencia, Transbank o cheque. Si la OT no está en un manifiesto suyo, responde 404. Si ya tiene resultado, responde 400. No puede pedir el listado general de órdenes, los clientes, los usuarios ni los reportes.")
    p(doc, "La cuenta de prueba es peoneta.prueba, con contraseña Peoneta2026!. El segundo paso del ingreso sigue enviando el código al correo ra.roa@duocuc.cl.")

    doc.add_heading("5. Matriz de permisos", 1)
    tabla(doc, ["Acción", "Administrador", "Operador", "Peoneta"], [
        ["Entrar a su propia vista", "Sí", "Sí", "Sí"],
        ["Entrar a la vista de otro rol", "No", "No", "No"],
        ["Ver todas las órdenes", "Sí", "Sí", "No"],
        ["Crear, editar y anular una orden", "Sí", "Sí", "No"],
        ["Cambiar el estado de una orden", "Sí", "Sí", "Solo entregada o rechazada, y solo de su manifiesto"],
        ["Crear manifiesto y sacarlo a ruta", "Sí", "Sí", "No"],
        ["Ver solo su manifiesto", "Si está asignado", "No aplica", "Sí"],
        ["Registrar entrega, comentario y pago", "Si está asignado", "No", "Sí"],
        ["Ver bodega y recibir bultos", "Sí", "Sí", "No"],
        ["Cobrar con Webpay", "Sí", "Sí", "No. En destino anota el método"],
        ["Ver clientes", "Sí", "Sí", "No"],
        ["Crear conductor o camión", "Sí", "No. Solo los consulta", "No"],
        ["Crear o editar tarifas", "Sí", "No. Solo las consulta", "No"],
        ["Cambiar roles", "Sí", "No", "No"],
        ["Ver reportes y exportarlos", "Sí", "No", "No"],
        ["Ver alertas", "Sí", "Sí", "No"],
    ])

    doc.add_heading("6. Lo que ningún rol puede hacer", 1)
    p(doc, "Nadie entra sin token, salvo el sitio público, el login, el registro y el retorno de Webpay. Nadie borra físicamente una entrega: la base lo impide para conservar la trazabilidad. Nadie elimina una orden de verdad: la anulación deja el registro. Nadie asigna el manifiesto de otra persona desde el cuerpo de la petición del peoneta. Y nadie queda como único administrador y luego se degrada a otro rol: el servicio lo rechaza.")
    p(doc, "El cliente, cuando se complete su vista, deberá ver solo las órdenes de la empresa ligada en cliente_usuario. Hoy esa vista no muestra datos de operación y tampoco puede cruzar al administrador, al operador ni al peoneta.")

    destino = BASE / "Informe-roles-Marexpress.docx"
    doc.save(destino)
    return destino


if __name__ == "__main__":
    for ruta in (informe_pruebas(), informe_codigos(), informe_roles()):
        print(ruta.name, ruta.stat().st_size)
